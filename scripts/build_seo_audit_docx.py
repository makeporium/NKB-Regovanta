from pathlib import Path
import re
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


def set_font(run, name="Arial", size=None, bold=None, color="000000"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    set_font(run, size=8, color="666666")
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instruction)
    run._r.append(end)


def normalize_markdown(text):
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", text)
    text = text.replace("`", "")
    return text


def add_mixed_text(paragraph, text, size=9):
    text = normalize_markdown(text)
    parts = re.split(r"(\*\*.*?\*\*)", text)
    for part in parts:
        if not part:
            continue
        bold = part.startswith("**") and part.endswith("**")
        value = part[2:-2] if bold else part
        run = paragraph.add_run(value)
        set_font(run, size=size, bold=bold)


def configure_document(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.62)
    section.bottom_margin = Inches(0.62)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    section.header_distance = Inches(0.25)
    section.footer_distance = Inches(0.25)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08

    title = styles["Title"]
    title.font.name = "Arial"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    title.font.size = Pt(25)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(14)

    for name, size, before, after in (
        ("Heading 1", 16, 14, 7),
        ("Heading 2", 11.5, 8, 3),
        ("Heading 3", 10.5, 7, 3),
    ):
        style = styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    if "Audit Detail" not in styles:
        detail = styles.add_style("Audit Detail", 1)
    else:
        detail = styles["Audit Detail"]
    detail.font.name = "Arial"
    detail._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    detail._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    detail.font.size = Pt(8.4)
    detail.font.color.rgb = RGBColor(0, 0, 0)
    detail.paragraph_format.left_indent = Inches(0.12)
    detail.paragraph_format.first_line_indent = Inches(-0.12)
    detail.paragraph_format.space_after = Pt(1.5)
    detail.paragraph_format.line_spacing = 1.0

    header = section.header.paragraphs[0]
    header.text = "NKB Regovanta Page by Page SEO Audit"
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in header.runs:
        set_font(run, size=8, color="666666")
    add_page_number(section.footer.paragraphs[0])


def add_summary_table(doc, rows):
    table = doc.add_table(rows=1, cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False
    widths = [Inches(5.8), Inches(1.05)] if len(rows[0]) == 2 else None
    for index, value in enumerate(rows[0]):
        cell = table.rows[0].cells[index]
        cell.text = value
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, "17365D")
        set_cell_margins(cell)
        if widths:
            cell.width = widths[index]
        for run in cell.paragraphs[0].runs:
            set_font(run, size=8.5, bold=True, color="FFFFFF")
    for row_index, values in enumerate(rows[1:]):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cell = cells[index]
            cell.text = normalize_markdown(value)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if widths:
                cell.width = widths[index]
            if row_index % 2:
                set_cell_shading(cell, "F3F6FA")
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                if index == len(values) - 1:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    set_font(run, size=8.3)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def build(markdown_path, output_path):
    lines = markdown_path.read_text(encoding="utf-8-sig").splitlines()
    doc = Document()
    configure_document(doc)
    in_inventory = False
    pending_table = []
    first_inventory_group = True

    def flush_table():
        nonlocal pending_table
        if pending_table:
            add_summary_table(doc, pending_table)
            pending_table = []

    for raw in lines:
        line = raw.rstrip()
        if line.startswith("|"):
            cells = [cell.strip() for cell in line.strip("|").split("|")]
            if all(re.fullmatch(r":?-+:?", cell) for cell in cells):
                continue
            pending_table.append(cells)
            continue
        flush_table()
        if not line:
            continue
        if line.startswith("# "):
            paragraph = doc.add_paragraph(style="Title")
            add_mixed_text(paragraph, line[2:], size=25)
            subtitle = doc.add_paragraph("Technical eligibility, Search Console evidence, and recommended decisions for every route")
            subtitle.paragraph_format.space_after = Pt(12)
            for run in subtitle.runs:
                set_font(run, size=11, color="444444")
        elif line.startswith("## "):
            heading = line[3:]
            force_page_break = False
            if heading == "Page by page inventory":
                force_page_break = True
                in_inventory = True
            elif in_inventory:
                if not first_inventory_group:
                    force_page_break = True
                first_inventory_group = False
            paragraph = doc.add_heading(normalize_markdown(heading), level=1)
            if force_page_break:
                paragraph.paragraph_format.page_break_before = True
        elif line.startswith("### "):
            paragraph = doc.add_heading(normalize_markdown(line[4:]), level=2)
            paragraph.paragraph_format.keep_with_next = True
        elif re.match(r"^\d+\. ", line):
            # Preserve the explicit Markdown number so each independent list
            # starts where the report author intended instead of continuing a
            # previous Word auto-numbered list.
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.22)
            paragraph.paragraph_format.first_line_indent = Inches(-0.18)
            add_mixed_text(paragraph, line, size=9.2)
        elif line.startswith("- "):
            style = "Audit Detail" if in_inventory else "List Bullet"
            paragraph = doc.add_paragraph(style=style)
            add_mixed_text(paragraph, line[2:], size=8.4 if in_inventory else 9.2)
            if in_inventory:
                paragraph.paragraph_format.keep_with_next = not line.startswith("- **Recommended decision:")
                if line.startswith("- **Recommended decision:"):
                    paragraph.paragraph_format.space_after = Pt(5)
        else:
            paragraph = doc.add_paragraph()
            add_mixed_text(paragraph, line, size=9.5)
    flush_table()

    core = doc.core_properties
    core.title = "NKB Regovanta Page by Page SEO Audit"
    core.subject = "Technical SEO and Google Search Console review of all website routes"
    core.author = "NKB Regovanta"
    core.keywords = "SEO audit, Google Search Console, CDSCO, EU MDR, EU IVDR, US FDA"
    doc.save(output_path)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit("Usage: build_seo_audit_docx.py input.md output.docx")
    build(Path(sys.argv[1]), Path(sys.argv[2]))
