"""Apply the approved URL Map sheet to the TanStack route tree and redirects."""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import subprocess
import time
from pathlib import Path

import openpyxl


PROJECT = Path(__file__).resolve().parents[1]
WORKBOOK = Path(r"C:\Users\astro\Desktop\AYUSH ALL\Client Projects\OTHER WORK\NKB_Regovanta_URL_Slug_Map.xlsx")
TEXT_EXTENSIONS = {".ts", ".tsx", ".js", ".jsx", ".cjs", ".mjs", ".json", ".xml", ".txt", ".md", ".csv"}
SKIP_NAMES = {"routeTree.gen.ts", "apply_excel_url_slug_map.py"}


def write_text(path: Path, text: str) -> None:
    last_error = None
    for _ in range(5):
        try:
            with open(path, "w", encoding="utf-8", newline="") as handle:
                handle.write(text)
            return
        except OSError as error:
            last_error = error
            time.sleep(0.15)
    raise last_error


def load_mapping() -> list[dict[str, str]]:
    workbook = openpyxl.load_workbook(WORKBOOK, read_only=True, data_only=True)
    sheet = workbook["URL Map"]
    rows = sheet.iter_rows(values_only=True)
    headers = next(rows)
    records = [dict(zip(headers, row)) for row in rows]
    mapping = []
    for row in records:
        if row.get("Action") != "RENAME":
            continue
        mapping.append(
            {
                "old_route": row["CurrentURL"],
                "new_route": row["FinalURL"],
                "old_file": row["CurrentFile"],
                "new_file": row["NewFile"],
                "indexability": row["Indexability"],
                "priority": row["Priority"],
            }
        )
    return mapping


def validate(mapping: list[dict[str, str]]) -> None:
    errors: list[str] = []
    old_routes = [m["old_route"] for m in mapping]
    new_routes = [m["new_route"] for m in mapping]
    old_files = [m["old_file"] for m in mapping]
    new_files = [m["new_file"] for m in mapping]

    for label, values in (("old route", old_routes), ("new route", new_routes), ("old file", old_files), ("new file", new_files)):
        duplicates = sorted({value for value in values if values.count(value) > 1})
        if duplicates:
            errors.append(f"Duplicate {label}s: {duplicates}")

    untouched_routes = set()
    workbook = openpyxl.load_workbook(WORKBOOK, read_only=True, data_only=True)
    sheet = workbook["URL Map"]
    rows = sheet.iter_rows(values_only=True)
    headers = next(rows)
    for row in rows:
        record = dict(zip(headers, row))
        if record.get("Action") == "KEEP":
            untouched_routes.add(record.get("FinalURL"))
    collisions = sorted(set(new_routes) & untouched_routes)
    if collisions:
        errors.append(f"New routes collide with KEEP routes: {collisions}")

    for item in mapping:
        old_path = PROJECT / item["old_file"]
        new_path = PROJECT / item["new_file"]
        if not old_path.is_file():
            errors.append(f"Missing source file: {item['old_file']}")
            continue
        if new_path.exists() and new_path != old_path:
            errors.append(f"Destination already exists: {item['new_file']}")
        text = old_path.read_text(encoding="utf-8")
        if item["old_route"] not in text:
            errors.append(f"Route string not found in source file: {item['old_file']} -> {item['old_route']}")

    if errors:
        raise SystemExit("Validation failed:\n- " + "\n- ".join(errors))


def iter_text_files():
    for root_name in ("src", "public", "scripts", "supabase"):
        root = PROJECT / root_name
        if not root.exists():
            continue
        for path in root.rglob("*"):
            if path.is_file() and path.suffix.lower() in TEXT_EXTENSIONS and path.name not in SKIP_NAMES:
                yield path
    for name in ("vite.config.ts", "vercel.json", ".noindex-review-evidence.json"):
        path = PROJECT / name
        if path.is_file() and path.name != "vercel.json":
            yield path


def simultaneous_replace(text: str, mapping: list[dict[str, str]]) -> tuple[str, int]:
    replacements = {item["old_route"]: item["new_route"] for item in mapping}
    pattern = re.compile("|".join(re.escape(route) for route in sorted(replacements, key=len, reverse=True)))
    count = 0

    def substitute(match: re.Match[str]) -> str:
        nonlocal count
        count += 1
        return replacements[match.group(0)]

    return pattern.sub(substitute, text), count


def replace_routes(mapping: list[dict[str, str]]) -> tuple[int, int]:
    changed_files = 0
    replacements = 0
    for path in iter_text_files():
        text = path.read_text(encoding="utf-8")
        updated, file_replacements = simultaneous_replace(text, mapping)
        if updated != text:
            write_text(path, updated)
            changed_files += 1
            replacements += file_replacements
    return changed_files, replacements


def git_head_text(relative_path: str) -> str:
    return subprocess.check_output(
        ["git", "show", f"HEAD:{relative_path.replace(os.sep, '/')}"],
        cwd=PROJECT,
    ).decode("utf-8")


def repair_from_head(mapping: list[dict[str, str]]) -> tuple[int, int]:
    """Rebuild migration-touched files from HEAD and apply non-cascading replacements."""
    modified = subprocess.check_output(
        ["git", "-c", "core.quotepath=false", "diff", "--name-only", "--diff-filter=M", "-z"],
        cwd=PROJECT,
    ).decode("utf-8").rstrip("\0").split("\0")
    restored = 0
    replacements = 0
    for relative in modified:
        normalized = relative.replace("\\", "/")
        if normalized in {"seo_audit_data.json", "vercel.json", "src/routeTree.gen.ts"}:
            continue
        if not normalized.startswith(("src/", "public/", "scripts/", "supabase/")) and normalized != "vite.config.ts":
            continue
        path = PROJECT / relative
        original = git_head_text(normalized)
        updated, count = simultaneous_replace(original, mapping)
        write_text(path, updated)
        restored += 1
        replacements += count

    # Every renamed route is rebuilt directly from its original tracked source.
    for item in mapping:
        original = git_head_text(item["old_file"])
        updated, count = simultaneous_replace(original, mapping)
        path = PROJECT / item["new_file"]
        write_text(path, updated)
        replacements += count

    updated_declarations = update_route_declarations(mapping)
    return restored + len(mapping) + updated_declarations, replacements


def rename_route_files(mapping: list[dict[str, str]]) -> int:
    renamed = 0
    for item in mapping:
        old_path = PROJECT / item["old_file"]
        new_path = PROJECT / item["new_file"]
        new_path.parent.mkdir(parents=True, exist_ok=True)
        os.replace(old_path, new_path)
        renamed += 1
    return renamed


def update_route_declarations(mapping: list[dict[str, str]]) -> int:
    """Update createFileRoute IDs while preserving TanStack pathless underscore segments."""
    updated_count = 0
    pattern = re.compile(r'(createFileRoute\(["\'])([^"\']+)(["\']\))')
    for item in mapping:
        path = PROJECT / item["new_file"]
        text = path.read_text(encoding="utf-8")
        match = pattern.search(text)
        if not match:
            raise RuntimeError(f"Missing createFileRoute declaration: {item['new_file']}")
        current_literal = match.group(2)
        normalized = current_literal.replace("_", "")
        if normalized == item["new_route"]:
            continue
        if normalized != item["old_route"]:
            raise RuntimeError(
                f"Unexpected route declaration in {item['new_file']}: {current_literal} "
                f"(expected public route {item['old_route']})"
            )
        literal_parent = current_literal.rsplit("/", 1)[0]
        new_leaf = item["new_route"].rsplit("/", 1)[-1]
        new_literal = f"{literal_parent}/{new_leaf}"
        updated = text[:match.start(2)] + new_literal + text[match.end(2):]
        write_text(path, updated)
        updated_count += 1
    return updated_count


def update_vercel_redirects(mapping: list[dict[str, str]]) -> int:
    path = PROJECT / "vercel.json"
    config = json.loads(path.read_text(encoding="utf-8"))
    redirects = config.setdefault("redirects", [])
    migration = {item["old_route"]: item["new_route"] for item in mapping}

    # Existing historical redirects should land directly on the newest URL.
    for redirect in redirects:
        seen = set()
        destination = redirect.get("destination")
        while destination in migration and destination not in seen:
            seen.add(destination)
            destination = migration[destination]
        redirect["destination"] = destination

    by_source = {redirect.get("source"): redirect for redirect in redirects}
    for item in mapping:
        by_source[item["old_route"]] = {
            "source": item["old_route"],
            "destination": item["new_route"],
            "permanent": True,
        }
    config["redirects"] = list(by_source.values())
    path.write_text(json.dumps(config, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    return len(mapping)


def update_review_manifests(mapping: list[dict[str, str]]) -> None:
    route_map = {item["old_route"]: item["new_route"] for item in mapping}
    file_map = {
        Path(item["old_file"]).name: Path(item["new_file"]).name
        for item in mapping
    }

    review_path = PROJECT / "src/content/indexing-review.json"
    review = json.loads(review_path.read_text(encoding="utf-8"))
    for item in review:
        if item.get("route") in route_map:
            item["route"] = route_map[item["route"]]
        if item.get("relatedPage") in route_map:
            item["relatedPage"] = route_map[item["relatedPage"]]
        if item.get("file") in file_map:
            item["file"] = file_map[item["file"]]
    write_text(review_path, json.dumps(review, indent=2, ensure_ascii=False) + "\n")

    evidence_path = PROJECT / ".noindex-review-evidence.json"
    evidence_text = evidence_path.read_text(encoding="utf-8")
    for old_name, new_name in file_map.items():
        evidence_text = evidence_text.replace(old_name, new_name)
    write_text(evidence_path, evidence_text)

    replacements_path = PROJECT / "src/content/article-replacements.json"
    replacements = json.loads(replacements_path.read_text(encoding="utf-8"))
    for item in replacements:
        article_path = PROJECT / "src/content/articles" / f"{item['slug']}.json"
        normalized = article_path.read_text(encoding="utf-8").replace("\r\n", "\n")
        item["contentSha256"] = hashlib.sha256(normalized.encode("utf-8")).hexdigest()
    write_text(replacements_path, json.dumps(replacements, indent=2, ensure_ascii=False) + "\n")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--apply", action="store_true")
    parser.add_argument("--repair-from-head", action="store_true")
    args = parser.parse_args()

    mapping = load_mapping()
    if args.repair_from_head:
        files, replacements = repair_from_head(mapping)
        update_vercel_redirects(mapping)
        print(f"Rebuilt/repaired {files} files with {replacements} non-cascading URL replacements.")
        return
    validate(mapping)
    indexable = sum(item["indexability"] == "INDEXABLE" for item in mapping)
    noindex = len(mapping) - indexable
    print(f"Validated {len(mapping)} renames ({indexable} indexable, {noindex} noindex).")
    if not args.apply:
        print("Dry run complete. Re-run with --apply to make changes.")
        return

    renamed = rename_route_files(mapping)
    changed_files, replacements = replace_routes(mapping)
    redirects = update_vercel_redirects(mapping)
    print(f"Renamed {renamed} route files.")
    print(f"Updated {replacements} URL references across {changed_files} files.")
    print(f"Added or refreshed {redirects} permanent redirects.")


if __name__ == "__main__":
    main()
